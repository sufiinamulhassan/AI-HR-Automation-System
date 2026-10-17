import base64
import io
import logging
import re
from datetime import datetime, timezone

from config.database import get_db
from config.settings import settings
from config.tenancy import current_company
from shared.utils import mongo_doc as _mongo_doc


def mongo_doc(doc: dict) -> dict:
    out = _mongo_doc(doc)
    out.pop("logo_base64", None)
    return out

logger = logging.getLogger(__name__)

SINGLETON_ID = "singleton"

MAX_LETTERHEAD_DOCX_BYTES = 5 * 1024 * 1024

DEFAULT_PRIMARY_COLOR = "#4778f3"
DEFAULT_ACCENT_COLOR = "#222022"

_STALE_PRIMARY_COLOR = "#fca311"
_STALE_ACCENT_COLOR = "#14213d"

_EDITABLE_FIELDS = (
    "company_logo_url", "primary_color", "accent_color",
    "letterhead_header_text", "letterhead_footer_text",
)
_HEX_COLOR_RE = re.compile(r"^#[0-9a-fA-F]{6}$")

LOGO_PATH_PREFIX = "/api/v1/branding/logo/"
MAX_LOGO_BYTES = 1_000_000
_LOGO_FIELDS = ("logo_base64", "logo_content_type", "logo_sha")


def sniff_logo_type(data: bytes) -> str | None:
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    return None


def is_uploaded_logo_url(value) -> bool:
    return isinstance(value, str) and value.startswith(LOGO_PATH_PREFIX)


def _heal_stale_defaults(doc: dict) -> dict:
    fix = {}
    if doc.get("primary_color") == _STALE_PRIMARY_COLOR:
        fix["primary_color"] = DEFAULT_PRIMARY_COLOR
    if doc.get("accent_color") == _STALE_ACCENT_COLOR:
        fix["accent_color"] = DEFAULT_ACCENT_COLOR
    if fix:
        fix["updated_at"] = datetime.now(timezone.utc).isoformat()
    return fix


def _default_doc() -> dict:
    now = datetime.now(timezone.utc).isoformat()
    return {
        "_id": SINGLETON_ID,
        "company_logo_url": "",
        "primary_color": DEFAULT_PRIMARY_COLOR,
        "accent_color": DEFAULT_ACCENT_COLOR,
        "letterhead_header_text": "",
        "letterhead_footer_text": "",
        "letterhead_docx_base64": None,
        "letterhead_docx_filename": None,
        "letterhead_template": None,
        "created_at": now,
        "updated_at": now,
    }


async def get_branding_settings() -> dict:
    db = get_db()
    if db is None:
        logger.warning("get_branding_settings: database unavailable, returning in-memory defaults")
        return mongo_doc(_default_doc())

    doc = await db.branding_settings.find_one({"_id": SINGLETON_ID})
    if doc is not None:
        healed = _heal_stale_defaults(doc)
        if healed:
            await db.branding_settings.update_one({"_id": SINGLETON_ID}, {"$set": healed})
            doc.update(healed)
        return mongo_doc(doc)

    fresh = _default_doc()
    try:
        await db.branding_settings.insert_one(dict(fresh))
    except Exception as exc:
        logger.info("branding_settings singleton insert race, re-reading: %s", exc)
        existing = await db.branding_settings.find_one({"_id": SINGLETON_ID})
        if existing is not None:
            return mongo_doc(existing)
    return mongo_doc(fresh)


def product_default_branding() -> dict:
    return {
        "company_logo_url": "",
        "primary_color": DEFAULT_PRIMARY_COLOR,
        "accent_color": DEFAULT_ACCENT_COLOR,
        "company_name": "",
        "company_slug": "",
        "terms_url": settings.TERMS_URL or "",
        "privacy_url": settings.PRIVACY_URL or "",
    }


async def _legal_links() -> dict:
    from services.hr_module.company_settings_service import SINGLETON_ID as PROFILE_ID, clean_legal_url

    stored: dict = {}
    db = get_db()
    if db is not None:
        stored = await db.company_settings.find_one({"_id": PROFILE_ID}, {"terms_url": 1, "privacy_url": 1}) or {}
    out = {}
    for key, platform_default in (("terms_url", settings.TERMS_URL), ("privacy_url", settings.PRIVACY_URL)):
        try:
            out[key] = clean_legal_url(stored.get(key)) or platform_default or ""
        except ValueError:
            out[key] = platform_default or ""
    return out


async def get_public_branding(company_resolved: bool = True) -> dict:
    if not company_resolved:
        return product_default_branding()
    from services.hr_module.company_settings_service import display_name

    doc = await get_branding_settings()
    return {
        "company_logo_url": doc.get("company_logo_url") or "",
        "primary_color": doc.get("primary_color") or DEFAULT_PRIMARY_COLOR,
        "accent_color": doc.get("accent_color") or DEFAULT_ACCENT_COLOR,
        "company_name": await display_name(),
        "company_slug": current_company().get("slug") or "",
        **await _legal_links(),
    }


def _validate_updates(updates: dict) -> dict:
    clean = {k: v for k, v in updates.items() if k in _EDITABLE_FIELDS and v is not None}
    for field in ("primary_color", "accent_color"):
        if field in clean and clean[field] != "" and not _HEX_COLOR_RE.match(clean[field]):
            raise ValueError(f"{field} must be a 6-digit hex color, e.g. #4778f3")
    return clean


async def update_branding_settings(updates: dict, updated_by: str | None = None) -> dict:
    clean = _validate_updates(updates)

    db = get_db()
    if db is None:
        logger.warning("update_branding_settings: database unavailable, update not persisted")
        merged = _default_doc()
        merged.update(clean)
        return mongo_doc(merged)

    await get_branding_settings()

    clean["updated_at"] = datetime.now(timezone.utc).isoformat()
    if updated_by:
        clean["updated_by"] = updated_by

    change: dict = {"$set": clean}
    if "company_logo_url" in clean and not is_uploaded_logo_url(clean["company_logo_url"]):
        change["$unset"] = {field: "" for field in _LOGO_FIELDS}
    await db.branding_settings.update_one({"_id": SINGLETON_ID}, change, upsert=True)
    doc = await db.branding_settings.find_one({"_id": SINGLETON_ID})
    return mongo_doc(doc) if doc is not None else mongo_doc(_default_doc())


async def store_logo(data: bytes, content_type: str, *, company_id: str, updated_by: str | None = None) -> dict:
    import base64
    import hashlib

    db = get_db()
    if db is None:
        raise RuntimeError("Database unavailable")
    await get_branding_settings()
    sha = hashlib.sha256(data).hexdigest()
    now = datetime.now(timezone.utc).isoformat()
    fields = {
        "logo_base64": base64.b64encode(data).decode("ascii"),
        "logo_content_type": content_type,
        "logo_sha": sha,
        "company_logo_url": f"{LOGO_PATH_PREFIX}{company_id}?v={sha[:12]}",
        "updated_at": now,
    }
    if updated_by:
        fields["updated_by"] = updated_by
    await db.branding_settings.update_one({"_id": SINGLETON_ID}, {"$set": fields}, upsert=True)
    return mongo_doc(await db.branding_settings.find_one({"_id": SINGLETON_ID}))


async def load_logo() -> tuple[bytes, str, str] | None:
    import base64

    db = get_db()
    if db is None:
        return None
    doc = await db.branding_settings.find_one(
        {"_id": SINGLETON_ID}, {"logo_base64": 1, "logo_content_type": 1, "logo_sha": 1},
    )
    if not doc or not doc.get("logo_base64"):
        return None
    try:
        data = base64.b64decode(doc["logo_base64"])
    except ValueError:
        return None
    return data, doc.get("logo_content_type") or "image/png", doc.get("logo_sha") or ""


async def resolve_logo_bytes(branding: dict | None) -> bytes | None:
    branding = branding or {}
    template_logo = (branding.get("letterhead_template") or {}).get("logo_base64")
    if template_logo:
        try:
            return base64.b64decode(template_logo)
        except Exception as exc:
            logger.warning("letterhead template logo decode failed, rendering without it | error=%s", exc)
            return None
    logo_url = branding.get("company_logo_url")
    if not logo_url:
        return None
    try:
        if is_uploaded_logo_url(logo_url):
            stored = await load_logo()
            if stored is None:
                raise ValueError("the uploaded logo is missing")
            return stored[0]
        from services.hr_module.url_safety import fetch_public_image

        return await fetch_public_image(logo_url)
    except Exception as exc:
        logger.warning("letterhead logo fetch failed, rendering without it | error=%s", exc)
        return None


async def clear_logo(*, updated_by: str | None = None) -> dict:
    db = get_db()
    if db is None:
        raise RuntimeError("Database unavailable")
    await get_branding_settings()
    fields = {"company_logo_url": "", "updated_at": datetime.now(timezone.utc).isoformat()}
    if updated_by:
        fields["updated_by"] = updated_by
    await db.branding_settings.update_one(
        {"_id": SINGLETON_ID}, {"$set": fields, "$unset": {field: "" for field in _LOGO_FIELDS}},
    )
    return mongo_doc(await db.branding_settings.find_one({"_id": SINGLETON_ID}))


def parse_letterhead_docx(raw_bytes: bytes) -> dict:
    from docx import Document

    try:
        doc = Document(io.BytesIO(raw_bytes))
    except Exception as exc:
        raise ValueError(f"Could not read DOCX file: {exc}")

    header_parts: list[str] = []
    footer_parts: list[str] = []
    header_footer_parts = []
    for section in doc.sections:
        try:
            header_parts += [p.text for p in section.header.paragraphs if p.text.strip()]
            header_footer_parts.append(section.header.part)
        except Exception as exc:
            logger.warning("parse_letterhead_docx: could not read a section header: %s", exc)
        try:
            footer_parts += [p.text for p in section.footer.paragraphs if p.text.strip()]
            header_footer_parts.append(section.footer.part)
        except Exception as exc:
            logger.warning("parse_letterhead_docx: could not read a section footer: %s", exc)

    body_paragraphs = [p.text for p in doc.paragraphs if p.text.strip() and "{{" in p.text]

    logo_base64 = None
    logo_content_type = None
    for part in [doc.part, *header_footer_parts]:
        try:
            rels = part.rels
        except Exception:
            continue
        for rel in rels.values():
            if "image" not in getattr(rel, "reltype", ""):
                continue
            try:
                blob = rel.target_part.blob
            except Exception:
                continue
            logo_base64 = base64.b64encode(blob).decode("ascii")
            logo_content_type = getattr(rel.target_part, "content_type", None) or "image/png"
            break
        if logo_base64:
            break

    return {
        "header_text": "\n".join(header_parts),
        "footer_text": "\n".join(footer_parts),
        "logo_base64": logo_base64,
        "logo_content_type": logo_content_type,
        "body_paragraphs": body_paragraphs,
    }


async def set_letterhead_docx_template(raw_bytes: bytes, filename: str | None, updated_by: str | None = None) -> dict:
    if len(raw_bytes) > MAX_LETTERHEAD_DOCX_BYTES:
        raise ValueError(
            f"Letterhead template file is too large (max {MAX_LETTERHEAD_DOCX_BYTES // (1024 * 1024)}MB)"
        )

    template = parse_letterhead_docx(raw_bytes)

    update = {
        "letterhead_docx_base64": base64.b64encode(raw_bytes).decode("ascii"),
        "letterhead_docx_filename": filename,
        "letterhead_template": template,
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    if updated_by:
        update["updated_by"] = updated_by

    db = get_db()
    if db is None:
        logger.warning("set_letterhead_docx_template: database unavailable, update not persisted")
        merged = _default_doc()
        merged.update(update)
        return mongo_doc(merged)

    await get_branding_settings()
    await db.branding_settings.update_one({"_id": SINGLETON_ID}, {"$set": update}, upsert=True)
    doc = await db.branding_settings.find_one({"_id": SINGLETON_ID})
    return mongo_doc(doc) if doc is not None else mongo_doc(_default_doc())


def get_letterhead_template_preview(branding: dict) -> dict:
    branding = branding or {}
    template = branding.get("letterhead_template") or {}
    return {
        "filename": branding.get("letterhead_docx_filename"),
        "header_text": template.get("header_text") or "",
        "footer_text": template.get("footer_text") or "",
        "has_logo": bool(template.get("logo_base64")),
        "body_paragraph_count": len(template.get("body_paragraphs") or []),
    }
