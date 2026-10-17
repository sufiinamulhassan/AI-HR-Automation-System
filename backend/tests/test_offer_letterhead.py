import io

import pytest

pytestmark = pytest.mark.anyio


def _build_docx(*, header_text=None, footer_text=None, body_text=None, with_logo=False) -> bytes:
    from docx import Document

    doc = Document()
    section = doc.sections[0]
    if header_text is not None:
        section.header.paragraphs[0].text = header_text
    if footer_text is not None:
        section.footer.paragraphs[0].text = footer_text
    if body_text is not None:
        doc.add_paragraph(body_text)
    if with_logo:
        from PIL import Image as PILImage

        buf = io.BytesIO()
        PILImage.new("RGB", (8, 8), color="green").save(buf, format="PNG")
        buf.seek(0)
        doc.add_picture(buf)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


async def test_parse_letterhead_docx_extracts_header_footer_logo_and_placeholders():
    from services.hr_module.branding_service import parse_letterhead_docx

    raw = _build_docx(
        header_text="Acme Corp, 1 Main St",
        footer_text="Confidential, Acme Corp",
        body_text="Dear {{candidate_name}}, welcome to {{company_name}} as {{job_title}}.",
        with_logo=True,
    )
    result = parse_letterhead_docx(raw)
    assert result["header_text"] == "Acme Corp, 1 Main St"
    assert "Confidential" in result["footer_text"]
    assert result["logo_base64"]
    assert result["logo_content_type"] == "image/png"
    assert any("{{candidate_name}}" in p for p in result["body_paragraphs"])


async def test_parse_letterhead_docx_with_no_logo_or_placeholders():
    from services.hr_module.branding_service import parse_letterhead_docx

    raw = _build_docx(header_text="Plain Header", footer_text="Plain Footer")
    result = parse_letterhead_docx(raw)
    assert result["header_text"] == "Plain Header"
    assert result["footer_text"] == "Plain Footer"
    assert result["logo_base64"] is None
    assert result["body_paragraphs"] == []


async def test_parse_letterhead_docx_rejects_garbage_bytes():
    from services.hr_module.branding_service import parse_letterhead_docx

    with pytest.raises(ValueError):
        parse_letterhead_docx(b"this is not a docx file")


async def test_set_letterhead_docx_template_persists_extracted_fields():
    from services.hr_module.branding_service import get_branding_settings, set_letterhead_docx_template

    raw = _build_docx(header_text="Persisted Header", body_text="Hi {{candidate_name}}")
    saved = await set_letterhead_docx_template(raw, "letterhead.docx", updated_by="admin@test.com")
    assert saved["letterhead_docx_filename"] == "letterhead.docx"
    assert saved["letterhead_template"]["header_text"] == "Persisted Header"

    refetched = await get_branding_settings()
    assert refetched["letterhead_template"]["header_text"] == "Persisted Header"
    assert refetched["letterhead_docx_base64"]


async def test_set_letterhead_docx_template_rejects_oversized_file(monkeypatch):
    from services.hr_module import branding_service

    monkeypatch.setattr(branding_service, "MAX_LETTERHEAD_DOCX_BYTES", 10)
    raw = _build_docx(header_text="Too big")
    with pytest.raises(ValueError):
        await branding_service.set_letterhead_docx_template(raw, "big.docx")


async def test_get_letterhead_template_preview_never_leaks_raw_bytes():
    from services.hr_module.branding_service import get_letterhead_template_preview

    branding = {
        "letterhead_docx_filename": "letterhead.docx",
        "letterhead_docx_base64": "verylongbase64...",
        "letterhead_template": {
            "header_text": "H", "footer_text": "F",
            "logo_base64": "imgbytes...", "body_paragraphs": ["{{candidate_name}}"],
        },
    }
    preview = get_letterhead_template_preview(branding)
    assert preview == {
        "filename": "letterhead.docx", "header_text": "H", "footer_text": "F",
        "has_logo": True, "body_paragraph_count": 1,
    }
    assert "letterhead_docx_base64" not in preview
    assert "logo_base64" not in preview


async def test_upload_letterhead_template_requires_admin(client, standard_headers):
    raw = _build_docx(header_text="X")
    resp = await client.post(
        "/api/v1/offers/letterhead-template",
        files={"file": ("letterhead.docx", raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=standard_headers,
    )
    assert resp.status_code == 403


async def test_upload_letterhead_template_rejects_non_docx(client, admin_headers):
    resp = await client.post(
        "/api/v1/offers/letterhead-template",
        files={"file": ("letterhead.pdf", b"%PDF-1.4 not a docx", "application/pdf")},
        headers=admin_headers,
    )
    assert resp.status_code == 400


async def test_upload_letterhead_template_rejects_unparseable_docx(client, admin_headers):
    resp = await client.post(
        "/api/v1/offers/letterhead-template",
        files={"file": ("letterhead.docx", b"not really a docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=admin_headers,
    )
    assert resp.status_code == 400


async def test_upload_and_get_letterhead_template_roundtrip(client, admin_headers):
    raw = _build_docx(
        header_text="HTTP Header", footer_text="HTTP Footer",
        body_text="Hi {{candidate_name}}", with_logo=True,
    )
    upload_resp = await client.post(
        "/api/v1/offers/letterhead-template",
        files={"file": ("letterhead.docx", raw, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=admin_headers,
    )
    assert upload_resp.status_code == 200, upload_resp.text
    body = upload_resp.json()
    assert body["header_text"] == "HTTP Header"
    assert body["has_logo"] is True
    assert body["body_paragraph_count"] == 1

    get_resp = await client.get("/api/v1/offers/letterhead-template", headers=admin_headers)
    assert get_resp.status_code == 200
    assert get_resp.json()["header_text"] == "HTTP Header"
